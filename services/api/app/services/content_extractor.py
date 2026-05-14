from __future__ import annotations

import io
from pathlib import Path


def _extract_docx(file_bytes: bytes) -> str:
    from docx import Document  # python-docx  # noqa: PLC0415

    doc = Document(io.BytesIO(file_bytes))
    parts: list[str] = []
    for para in doc.paragraphs:
        text = para.text.strip()
        if text:
            parts.append(text)
    for table in doc.tables:
        for row in table.rows:
            row_cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if row_cells:
                parts.append(" | ".join(row_cells))
    return "\n".join(parts)


def extract_text_from_bytes(filename: str, file_bytes: bytes) -> str:
    """Return plain text from a Word / TXT / Markdown file."""
    ext = Path(filename).suffix.lower().lstrip(".")

    if ext in ("docx", "docm"):
        try:
            return _extract_docx(file_bytes)
        except Exception as exc:  # noqa: BLE001
            raise ValueError(f"Không thể đọc file DOCX: {exc}") from exc

    if ext == "doc":
        # python-docx cannot read the old binary .doc format.
        # Try anyway in case the file is actually .docx mis-named.
        try:
            return _extract_docx(file_bytes)
        except Exception:  # noqa: BLE001
            raise ValueError(
                "File .doc (định dạng Word cũ) chưa được hỗ trợ trực tiếp. "
                "Vui lòng mở file trong Word và lưu lại dưới dạng .docx, rồi tải lên lại."
            )

    if ext in ("txt", "md", "markdown"):
        try:
            return file_bytes.decode("utf-8")
        except UnicodeDecodeError:
            return file_bytes.decode("latin-1", errors="replace")

    raise ValueError(f"Định dạng file '.{ext}' không được hỗ trợ. Hãy dùng .docx, .txt hoặc .md.")
